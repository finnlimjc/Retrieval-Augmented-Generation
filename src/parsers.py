import email
import openpyxl
import pdfplumber
import re

from dataclasses import dataclass, field
from docx import Document as DocxDocument
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from src.config import IMAGES_DIR, ATTACHMENTS_DIR

@dataclass
class Element:
    '''Normalized input data carrying its own structural metadata, extracted from the various file types'''
    text: str
    element_type: str
    metadata: dict = field(default_factory=dict) # Additional metadata not included in the metadata.jsonl

def convert_str_to_filename_or_url(name:str) -> str:
    # Match one or more characters that is not in the square brackets and replace with _
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_")

def rows_to_markdown(rows:list[list]) -> str:
    """Serialize a table as Markdown so row-to-header association survives embedding."""
    cleaned = []
    for row in rows:
        cells = ["" if col is None else str(col).replace("\n", " ").strip() for col in row]
        if any(cells):
            cleaned.append(cells)
    
    if not cleaned:
        return ""
    
    # Formatting the table
    n_columns = max(len(r) for r in cleaned) # Find the maximum number of columns
    padded = [r + [""] * (n_columns - len(r)) for r in cleaned] # Pad each row to match the maximum number of columns
    col_names = "| " + " | ".join(padded[0]) + " |"
    line_below_col = "| " + " | ".join(["---"] * n_columns) + " |"
    lines = [col_names, line_below_col]
    for row in padded[1:]:
        lines.append("| " + " | ".join(row) + " |")
    
    return "\n".join(lines)

def extract_pdf(path:Path) -> list[Element]:
    """
    Extract narrative text, tables and embedded images from a PDF.
    
    Inputs:
        - path: PDF file to read.
    
    Outputs:
        - elements: one narrative element per page that has text, one element per table, and one image element per embedded image.
    """
    elements, page_texts = [], []
    with pdfplumber.open(path) as pdf:
        # For each page, extract the text and the table
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            page_texts.append(page_text)
            if page_text.strip():
                elements.append(Element(page_text, "narrative", {"page_number": page_number})) # Narrative is a placeholder for MVP
            
            for table_index, table in enumerate(page.find_tables()):
                markdown = rows_to_markdown(table.extract())
                if markdown:
                    metadata = {"page_number": page_number, "table_index": table_index}
                    elements.append(Element(markdown, "table", metadata))
    
    # Images are saved but left with empty text; the vision stage fills them in later
    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages, start=1):
        for image_index, image in enumerate(page.images):
            suffix = Path(image.name).suffix or ".png" # PdfReader returns image.jpg, fallback to .png if no suffix is found
            image_file_name = f"{convert_str_to_filename_or_url(path.stem)}_p{page_number}_{image_index}{suffix}" # filename_p1_1.jpg
            image_path = IMAGES_DIR/image_file_name
            image_path.write_bytes(image.data)
            
            page_has_text = bool(page_texts[page_number-1].strip())
            metadata = {"page_number": page_number, "image_path": str(image_path)}
            elements.append(Element("", "figure" if page_has_text else "scan_page", metadata)) # scan_page is a quick fix
    
    return elements

def extract_xlsx(path:Path) -> list[Element]:
    workbook = openpyxl.load_workbook(str(path), data_only=True)
    elements = []
    for worksheet in workbook.worksheets:
        rows = [list(r) for r in worksheet.iter_rows(values_only=True)]
        markdown = rows_to_markdown(rows)
        if not markdown:
            continue
        
        # Future work: Handle Images
        is_hidden = worksheet.sheet_state != "visible"
        metadata = {"sheet_name": worksheet.title, "sheet_hidden": is_hidden}
        elements.append(Element(markdown, "hidden_sheet" if is_hidden else "sheet", metadata))
    
    return elements

def extract_md(path:Path) -> list[Element]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [Element(text, "narrative", {})] if text.strip() else [] # no additional metadata

def extract_docx(path:Path) -> list[Element]:
    document = DocxDocument(str(path))
    
    # Get Text
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    elements = []
    if paragraphs:
        elements.append(Element("\n".join(paragraphs), "narrative", {})) # future work: get the page number
    
    # Future work: Handle Images
    
    # Get Tables
    for table_index, table in enumerate(document.tables):
        markdown = rows_to_markdown([[col.text for col in row.cells] for row in table.rows])
        if markdown:
            elements.append(Element(markdown, "table", {"table_index": table_index}))
    return elements

def extract_pptx(path:Path) -> list[Element]:
    elements = []
    presentation_slides = Presentation(str(path)).slides
    
    for slide_number, slide in enumerate(presentation_slides, start=1):
        metadata = {"slide_number": slide_number}
        
        shape_texts = [s.text for s in slide.shapes if s.has_text_frame and s.text.strip()]
        slide_has_text = bool(shape_texts)
        if shape_texts:
            elements.append(Element("\n".join(shape_texts), "slide", metadata))
        
        # Tables: extracted the same way as the PDF pipeline, so a table doesn't
        # silently vanish just because it lives in a pptx instead of a pdf.
        for table_index, shape in enumerate(slide.shapes):
            if shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                markdown = rows_to_markdown(rows)
                if markdown:
                    table_metadata = {"slide_number": slide_number, "table_index": table_index}
                    elements.append(Element(markdown, "table", table_metadata))
        
        # Images: saved to disk with empty text, same convention as the PDF
        # pipeline's "figure"/"scan_page" split. A slide that is only a pasted
        # picture (no other text) is flagged as scan_slide
        for image_index, shape in enumerate(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                suffix = f".{image.ext}" if image.ext else ".png"
                image_file_name = f"{convert_str_to_filename_or_url(path.stem)}_s{slide_number}_{image_index}{suffix}"
                image_path = IMAGES_DIR / image_file_name
                image_path.write_bytes(image.blob)
                
                image_metadata = {"slide_number": slide_number, "image_path": str(image_path)}
                elements.append(Element("", "figure" if slide_has_text else "scan_slide", image_metadata))
        
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                elements.append(Element(notes, "speaker_notes", metadata))
    
    return elements

def email_body(message) -> str:
    parts = []
    for part in message.walk():
        # The email body will get duplicated without this checks
        if part.get_content_type() != "text/plain" or part.get_filename():
            continue
        
        payload = part.get_payload(decode=True)
        if payload:
            parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
    
    return "\n".join(parts)

def extract_eml(path:Path) -> list[Element]:
    with path.open(encoding="utf-8", errors="replace") as handle:
        message = email.message_from_file(handle)
    
    # Process the email headers and body
    email_fields = ("Subject", "From", "To", "Date")
    headers = [f"{field}: {message.get(field, '')}" for field in email_fields]
    elements = [Element("\n".join(headers) + "\n\n" + email_body(message), "email_body", {})]
    
    # Future work: Handle Images
    
    # Process attachments
    attachment_dir = ATTACHMENTS_DIR / convert_str_to_filename_or_url(path.stem)
    is_initial_attachment = True
    for part in message.walk():
        filename = part.get_filename()
        payload = part.get_payload(decode=True) if filename else None
        if not payload:
            continue
        
        # Create directory to store all attachments for this email
        if is_initial_attachment:
            attachment_dir.mkdir(parents=True, exist_ok=True)
            is_initial_attachment = False
        
        # Write attachment file
        attachment_path = attachment_dir / convert_str_to_filename_or_url(filename)
        attachment_path.write_bytes(payload)
        for element in extract(attachment_path):
            element.metadata.update({"is_attachment": True, "parent_path": str(path), "attachment_filename": filename})
            elements.append(element)
    
    return elements

def extract(path:Path) -> list[Element]:
    EXTRACTORS = {
        ".pdf": extract_pdf,
        ".docx": extract_docx,
        ".pptx": extract_pptx,
        ".xlsx": extract_xlsx,
        ".md": extract_md,
        ".eml": extract_eml
    }
    
    extractor = EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        raise ValueError(f"No extractor registered for {path.suffix!r} ({path})")
    return extractor(path)